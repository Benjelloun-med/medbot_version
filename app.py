from flask import Flask, render_template, request, redirect, send_file, url_for, flash, jsonify
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from flask_migrate import Migrate
import os
from dotenv import load_dotenv
from openai import OpenAI
from werkzeug.security import generate_password_hash, check_password_hash
import urllib.parse
from functools import wraps
from docx import Document
from datetime import datetime
import io
load_dotenv()

app = Flask(__name__)
app.secret_key = os.getenv('SECRET_KEY')
#app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///users.db'
#app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

#import openai

#openai.api_type = "azure"
#openai.api_base = os.getenv('AZURE_OPENAI_ENDPOINT')
#openai.api_version = os.getenv('AZURE_OPENAI_VERSION', '2023-05-15')
#openai.api_key = os.getenv('AZURE_OPENAI_KEY')
#AZURE_DEPLOYMENT = os.getenv('AZURE_OPENAI_DEPLOYMENT')


# Affichage des variables d'environnement pour le débogage
#print("Client ID:", os.getenv('GOOGLE_CLIENT_ID'))
#print("Client Secret:", os.getenv('GOOGLE_CLIENT_SECRET'))

# Configuration OpenAI
client = OpenAI(api_key=os.getenv('OPENAI_API_KEY'))

# Configuration OAuth avec plus de détails
"""oauth = OAuth(app)
google_config = {
    'name': 'google',
    'client_id': os.getenv('GOOGLE_CLIENT_ID'),
    'client_secret': os.getenv('GOOGLE_CLIENT_SECRET'),
    'api_base_url': 'https://www.googleapis.com/oauth2/v2/',
    'server_metadata_url': 'https://accounts.google.com/.well-known/openid-configuration',
    'client_kwargs': {
        'scope': 'openid email profile',
        'redirect_uri': 'http://127.0.0.1:5000/login/google/authorize'
    }
}
oauth.register(**google_config)
print("OAuth Config:", google_config)

oauth.register(
    name='github',
    client_id=os.getenv('GITHUB_CLIENT_ID'),
    client_secret=os.getenv('GITHUB_CLIENT_SECRET'),
    api_base_url='https://api.github.com/',
    authorize_url='https://github.com/login/oauth/authorize',
    access_token_url='https://github.com/login/oauth/access_token',
)"""

if odbc_str := os.getenv('ODBC_STR'):
    params = urllib.parse.quote_plus(odbc_str)
    app.config['SQLALCHEMY_DATABASE_URI'] = f"mssql+pyodbc:///?odbc_connect={params}"
#else:
    #app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///users.db'
#app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)
migrate = Migrate(app, db)
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'

class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    login = db.Column(db.String(100), unique=True)
    name = db.Column(db.String(100))
    password_hash = db.Column(db.String(128))
    is_admin = db.Column(db.Boolean, default=False)
    is_active = db.Column(db.Boolean, default=True)
    documents = db.relationship('GeneratedDocument', backref='user', lazy=True)

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)

    def check_password(self, password):
        return check_password_hash(self.password_hash, password)

class GeneratedDocument(db.Model):
    __tablename__ = 'document'
    id = db.Column(db.Integer, primary_key=True)
    filename = db.Column(db.String(255), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    content = db.Column(db.Text, nullable=False)


@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        login = request.form.get('login')  # Modification ici
        password = request.form.get('password')
        user = User.query.filter_by(login=login).first()  # Modification ici
        
        if user and user.check_password(password):
            if not user.is_active:
                flash('Ce compte a été désactivé. Veuillez contacter un administrateur.')
                return redirect(url_for('login'))
            login_user(user)
            return redirect(url_for('admin_dashboard' if user.is_admin else 'chatbot'))
        else:
            flash('Login ou mot de passe incorrect')

    return render_template('login.html')

"""@app.route('/register', methods=['GET', 'POST'])
def register():
    if current_user.is_authenticated:
        return redirect(url_for('chatbot'))
    if request.method == 'POST':
        login = request.form.get('login')
        name = request.form.get('name')
        password = request.form.get('password')
        
        if User.query.filter_by(login=login).first():
            flash('Ce login est déjà utilisé')
            return redirect(url_for('register'))
        
        user = User(login=login, name=name)
        user.set_password(password)
        db.session.add(user)
        db.session.commit()
        
        login_user(user)
        return redirect(url_for('chatbot'))
    
    return render_template('register.html')

@app.route('/login/google')
def google_login():
    return oauth.google.authorize_redirect(redirect_uri=url_for('google_authorize', _external=True))

@app.route('/login/google/authorize')
def google_authorize():
    token = oauth.google.authorize_access_token()
    resp = oauth.google.get('https://www.googleapis.com/oauth2/v2/userinfo', token=token)
    user_info = resp.json()
    
    user = User.query.filter_by(email=user_info.get('email')).first()
    if not user:
        user = User(
            email=user_info.get('email'),
            name=user_info.get('name', 'Unknown')
        )
        db.session.add(user)
        db.session.commit()
    
    login_user(user)
    return redirect(url_for('chatbot'))

@app.route('/login/github')
def github_login():
    return oauth.github.authorize_redirect(redirect_uri=url_for('github_authorize', _external=True))

@app.route('/login/github/authorize')
def github_authorize():
    token = oauth.github.authorize_access_token()
    resp = oauth.github.get('user', token=token)
    user_info = resp.json()
    
    user = User.query.filter_by(email=user_info['email']).first()
    if not user:
        user = User(
            email=user_info['email'],
            name=user_info['login']
        )
        db.session.add(user)
        db.session.commit()
    
    login_user(user)
    return redirect(url_for('chatbot'))
"""


@app.route('/chatbot')
@login_required
def chatbot():
    return render_template('chatbot.html')

def load_cdc_content():
    try:
        with open("cdc.txt", "r", encoding="utf-8") as file:
            return file.read()
    except Exception as e:
        print(f"Erreur lors de la lecture du fichier cdc.txt: {e}")
        return ""

@app.route('/api/chat', methods=['POST'])
@login_required
def chat():
    data = request.json
    user_message = data.get('message', '')
    print("Message reçu:", user_message)  # Debug
    
    try:
        print("Tentative d'appel à OpenAI...")  # Debug
        cdc_content = load_cdc_content()
        
        # Appel à l'API OpenAI avec la nouvelle syntaxe
        response = client.chat.completions.create(
            model="gpt-4",
            messages=[
                {"role": "system", "content": f"Vous êtes un assistant spécialisé dans la création de cahiers des charges. Votre rôle est de poser des questions un par un aux utilisateurs pour structurer et rédiger leurs cahiers des charges de manière professionnelle et l'envoyer sous forme word. Sans répondre aux questions hors de ce contexte. Basez-vous sur le contenu suivant pour les questions : {cdc_content}"},
                {"role": "user", "content": user_message}
            ],
            max_tokens=500
        )
        print("Réponse OpenAI reçue")  # Debug

        bot_response = response.choices[0].message.content
        print("Réponse:", bot_response)  # Debug
        return jsonify({"response": bot_response})
    except Exception as e:
        print("Erreur:", e)
        return jsonify({"error": str(e)}), 500

@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('index'))



# Fonction pour vérifier si l'utilisateur est admin
def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated or not current_user.is_admin:
            flash('Accès non autorisé')
            return redirect(url_for('index'))
        return f(*args, **kwargs)
    return decorated_function

# Routes d'administration
@app.route('/admin')
@login_required
@admin_required
def admin_dashboard():
    users = User.query.all()
    return render_template('admin/dashboard.html', users=users)

@app.route('/admin/user/<int:user_id>/edit', methods=['GET', 'POST'])
@login_required
@admin_required
def admin_edit_user(user_id):
    user = User.query.get_or_404(user_id)
    if request.method == 'POST':
        user.name = request.form.get('name')
        user.login = request.form.get('login')
        user.is_admin = 'is_admin' in request.form
        if request.form.get('password'):
            user.set_password(request.form.get('password'))
        db.session.commit()
        flash('Utilisateur mis à jour avec succès')
        return redirect(url_for('admin_dashboard'))
    return render_template('admin/edit_user.html', user=user)

@app.route('/admin/user/<int:user_id>/delete', methods=['POST'])
@login_required
@admin_required
def admin_delete_user(user_id):
    user = User.query.get_or_404(user_id)
    if user.id == current_user.id:
        flash('Vous ne pouvez pas supprimer votre propre compte')
        return redirect(url_for('admin_dashboard'))
    db.session.delete(user)
    db.session.commit()
    flash('Utilisateur supprimé avec succès')
    return redirect(url_for('admin_dashboard'))

@app.route('/admin/user/new', methods=['GET', 'POST'])
@login_required
@admin_required
def admin_new_user():
    if request.method != 'POST':
        return render_template('admin/new_user.html')
    login = request.form.get('login')
    name = request.form.get('name')
    password = request.form.get('password')
    is_admin = 'is_admin' in request.form

    # Vérification des champs obligatoires
    if not login or not name or not password:
        flash("Tous les champs sont obligatoires.", "danger")
        return redirect(url_for('admin_new_user'))

    # Vérification de l'unicité du login
    if User.query.filter_by(login=login).first():
        flash('Ce login est déjà utilisé', "danger")
        return redirect(url_for('admin_new_user'))

    # (Optionnel) Vérification de l'unicité du nom si nécessaire
    # if User.query.filter_by(name=name).first():
    #     flash('Ce nom est déjà utilisé', "danger")
    #     return redirect(url_for('admin_new_user'))

    user = User(login=login, name=name, is_admin=is_admin)
    user.set_password(password)
    db.session.add(user)
    db.session.commit()
    flash('Utilisateur créé avec succès', "success")
    return redirect(url_for('admin_dashboard'))


@app.route('/admin/user/<int:user_id>/toggle-status', methods=['POST'])
@login_required
@admin_required
def admin_toggle_user_status(user_id):
    user = User.query.get_or_404(user_id)
    if user.id == current_user.id:
        return jsonify({"error": "Vous ne pouvez pas modifier votre propre statut"}), 400
    
    user.is_active = not user.is_active
    db.session.commit()
    
    return jsonify({
        "success": True,
        "is_active": user.is_active,
        "message": f"Compte {'activé' if user.is_active else 'désactivé'} avec succès"
    })
    
@app.route('/api/save-word', methods=['POST'])
@login_required
def save_word():
    try:
        data = request.json
        messages = data.get('messages', [])

        # Vérifier qu'il y a des messages et récupérer le dernier
        if not messages or len(messages) == 0:
            return jsonify({"error": "Aucun message disponible à enregistrer."}), 400

        last_message = messages[-1].get('content', '')

        # Créer un document Word
        doc = Document()
        doc.add_heading('Cahier des Charges', 0)
        doc.add_paragraph(f'Date de création : {datetime.now().strftime("%d/%m/%Y")}')
        doc.add_paragraph(last_message)  # Ajouter UNIQUEMENT le dernier message IA

        # Sauvegarder le document dans un buffer
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        # Générer un nom de fichier
        filename = f'cahier_des_charges_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx'

        # Sauvegarder uniquement le dernier message dans la base de données
        document_record = GeneratedDocument(
            filename=filename,
            user_id=current_user.id,
            content=last_message  # Sauvegarde du dernier message seulement
        )
        db.session.add(document_record)
        db.session.commit()

        return send_file(
            file_stream,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            as_attachment=True,
            download_name=filename
        )
    except Exception as e:
        print("Erreur détaillée lors de la sauvegarde en Word :", e)
        return jsonify({"error": "Une erreur est survenue lors de la sauvegarde du document.", "details": str(e)}), 500
@app.route('/admin/documents')
@login_required
@admin_required
def admin_documents():
    documents = GeneratedDocument.query.order_by(GeneratedDocument.created_at.desc()).all()
    return render_template('admin/documents.html', documents=documents)

@app.route('/admin/document/<int:doc_id>/download')
@login_required
@admin_required
def admin_download_document(doc_id):
    document = GeneratedDocument.query.get_or_404(doc_id)
    
    # Créer un nouveau document Word
    doc = Document()
    doc.add_heading('Cahier des Charges', 0)
    doc.add_paragraph(f'Date de création : {document.created_at.strftime("%d/%m/%Y")}')
    doc.add_paragraph(document.content)
    
    # Sauvegarder le document dans un buffer
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    return send_file(
        file_stream,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        as_attachment=True,
        download_name=document.filename
    )

if __name__ == '__main__':
    with app.app_context():
        db.create_all()
    app.run(debug=True) 
